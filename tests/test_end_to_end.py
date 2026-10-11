import json

import mock_grafana
import pytest
from docx import Document

from gravekeeper.cli import main


@pytest.fixture()
def grafana(monkeypatch, tmp_path):
    srv, url = mock_grafana.start()
    monkeypatch.setenv("GRAFANA_TOKEN", "test-token")
    cfg = tmp_path / "gk.yaml"
    cfg.write_text(f"grafana:\n  url: {url}\nusage:\n  source: none\noutput_dir: {tmp_path / 'out'}\n")
    yield cfg, tmp_path / "out"
    srv.shutdown()


def _codes(result):
    return {d["uid"]: sorted(f["code"] for f in d["findings"]) for d in result["dashboards"]}


def test_scan_finds_every_zombie(grafana):
    cfg, out = grafana
    assert main(["scan", "-c", str(cfg)]) == 0
    result = json.loads((out / "findings.json").read_text())
    codes = _codes(result)

    assert codes["gk-node-overview"] == []
    assert codes["gk-billing"] == []
    assert codes["gk-node-copy2"] == ["C02", "C08"]
    assert "C08" in codes["gk-node-copy"]
    assert codes["gk-legacy"] == ["C07"]
    assert codes["gk-influx"] == ["C03"]
    assert codes["gk-prom-old"] == ["C04"]
    assert "C10" in codes["gk-empty-shell"] and "C09" not in codes["gk-empty-shell"]
    assert set(codes["gk-untitled"]) >= {"C02", "C10"}
    assert set(codes["gk-ex-employee"]) == {"C07", "C09"}

    skipped = {s["uid"]: s["reason"] for s in result["skipped"]}
    assert "Provisioned" in skipped["gk-provisioned"]
    assert "critical" in skipped["gk-payments"]

    by_uid = {d["uid"]: d for d in result["dashboards"]}
    assert by_uid["gk-node-overview"]["status"] == "Healthy"
    assert by_uid["gk-untitled"]["status"] == "Terminal"
    ds = {d["uid"]: d for d in result["datasources"]}
    assert ds["prometheus-old"]["ok"] is False


def test_view_counts_mark_abandoned(grafana, monkeypatch):
    cfg, out = grafana
    cfg.write_text(cfg.read_text().replace("source: none", "source: loki"))
    views = {"gk-node-overview": 5, "gk-loki-ingest": 3, "gk-billing": 2}
    monkeypatch.setattr("gravekeeper.cli.dashboard_views", lambda c, d: views)
    assert main(["scan", "-c", str(cfg)]) == 0
    result = json.loads((out / "findings.json").read_text())
    by_uid = {d["uid"]: d for d in result["dashboards"]}
    assert by_uid["gk-node-overview"]["status"] == "Healthy"
    assert "C01" in _codes(result)["gk-influx"]
    # The copy nobody opens is flagged, the original that people use is kept.
    assert "C08" not in _codes(result)["gk-node-overview"]
    assert by_uid["gk-untitled"]["status"] == "Deceased"
    assert by_uid["gk-node-copy2"]["status"] == "Terminal"


def test_report_and_certificates(grafana):
    cfg, out = grafana
    assert main(["run", "-c", str(cfg)]) == 0
    report = out / "estate-health-report.docx"
    assert report.exists()
    text = "\n".join(c.text for t in Document(report).tables for r in t.rows for c in r.cells)
    assert "Copy of Copy of Node Overview (old)" in text
    assert "{{" not in text and "{%" not in text
    certs = list((out / "certificates").glob("*.docx"))
    assert len(certs) >= 1
    cert_text = "\n".join(p.text for p in Document(certs[0]).paragraphs)
    assert "will be archived on" in cert_text


def test_bad_token_is_reported(grafana, monkeypatch):
    cfg, _ = grafana
    monkeypatch.setenv("GRAFANA_TOKEN", "wrong")
    assert main(["scan", "-c", str(cfg)]) == 2


def test_loki_down_gives_clean_error(grafana, capsys):
    cfg, _ = grafana
    cfg.write_text(cfg.read_text().replace("source: none", "source: loki\n  loki_url: http://127.0.0.1:1"))
    assert main(["scan", "-c", str(cfg)]) == 2
    assert "Cannot reach Loki" in capsys.readouterr().err
