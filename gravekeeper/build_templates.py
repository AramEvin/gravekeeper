"""Generate the DOCX templates (docxtpl / Jinja2 tags inside Word files).

Run `python -m gravekeeper.build_templates` to rebuild them. You can also open
the generated .docx files in Word, restyle them (fonts, colors, logo), and keep
the {{ ... }} / {% ... %} tags: Gravekeeper will render your version instead.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ORANGE = "F46800"
DARK = "1F2430"
GREY = "6B7280"
LIGHT = "F3F4F6"
FONT = "Calibri"
MONO = "Consolas"

OUT_DIR = Path(__file__).parent / "templates"


# ---------------------------------------------------------------- helpers
def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _borders(table, color: str = "D1D5DB") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _widths(table, widths_cm) -> None:
    """Fixed column widths that both Word and LibreOffice respect."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for i, col in enumerate(grid.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(int(Cm(widths_cm[i]).twips)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def _run(par, text, bold=False, italic=False, color=None, size=None, font=None):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _cell_text(cell, text, bold=False, color=None, size=10, font=None, italic=False):
    p = cell.paragraphs[0]
    _run(p, text, bold=bold, color=color, size=size, font=font, italic=italic)
    p.paragraph_format.space_after = Pt(0)
    return p


def _base_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(4)
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
        h.font.color.rgb = RGBColor.from_string(DARK)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(4)
    # footer
    fp = sec.footer.paragraphs[0]
    _run(fp, "Gravekeeper  ·  {{ tool }}  ·  scanned {{ scanned_at }}", color=GREY, size=8)
    return doc


def _banner(doc, kicker: str, title: str, subtitle: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    _no_borders(t)
    _widths(t, [17.0])
    c = t.rows[0].cells[0]
    _shade(c, DARK)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(10)
    _run(p, kicker, bold=True, color=ORANGE, size=9)
    p2 = c.add_paragraph()
    _run(p2, title, bold=True, color="FFFFFF", size=22)
    p3 = c.add_paragraph()
    p3.paragraph_format.space_after = Pt(10)
    _run(p3, subtitle, color="D1D5DB", size=10)
    doc.add_paragraph()


def _callout(doc, label: str, text: str, color: str) -> None:
    t = doc.add_table(rows=1, cols=2)
    _no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    _widths(t, [2.0, 15.0])
    left, right = t.rows[0].cells
    _shade(left, color)
    _shade(right, LIGHT)
    p = _cell_text(left, label, bold=True, color="FFFFFF", size=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cell_text(right, text, size=9.5)
    doc.add_paragraph()


def _loop_table(doc, headers, cells, loop: str, widths=None) -> None:
    """Header row + {%tr for %} + data row + {%tr endfor %}."""
    t = doc.add_table(rows=4, cols=len(headers))
    _borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        _shade(c, DARK)
        _cell_text(c, h, bold=True, color="FFFFFF", size=9)
    _cell_text(t.rows[1].cells[0], "{%tr for " + loop + " %}", size=8)
    for i, expr in enumerate(cells):
        _cell_text(t.rows[2].cells[i], expr, size=9, bold=expr.startswith("{% cellbg"))
    _cell_text(t.rows[3].cells[0], "{%tr endfor %}", size=8)
    _widths(t, widths or [17.0 / len(headers)] * len(headers))
    doc.add_paragraph()


def _kv_table(doc, rows: list[tuple[str, str]]) -> None:
    t = doc.add_table(rows=len(rows), cols=2)
    _borders(t)
    for i, (k, v) in enumerate(rows):
        a, b = t.rows[i].cells
        _shade(a, LIGHT)
        _cell_text(a, k, bold=True, size=9.5)
        if v.startswith("{% cellbg"):
            _cell_text(b, v, bold=True, size=9.5)
        else:
            _cell_text(b, v, size=9.5)
    _widths(t, [4.5, 12.5])
    doc.add_paragraph()


def _p(doc, text, **kw):
    p = doc.add_paragraph()
    _run(p, text, **kw)
    return p


def _code(doc, text: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    _no_borders(t)
    _widths(t, [17.0])
    c = t.rows[0].cells[0]
    _shade(c, "111827")
    _cell_text(c, text, color="E5E7EB", size=8.5, font=MONO)
    doc.add_paragraph()


# ---------------------------------------------------------------- templates
def build_estate_report(path: Path) -> None:
    doc = _base_document()
    _banner(
        doc,
        "GRAVEKEEPER  ·  ESTATE HEALTH REPORT",
        "Grafana Estate Health Report",
        "{{ grafana_url }}  ·  Grafana {{ grafana_version }}  ·  {{ scanned_date }}",
    )

    doc.add_heading("1. Executive summary", 1)
    _p(
        doc,
        "Gravekeeper scanned {{ total }} dashboards and {{ datasources|length }} datasources. "
        "{{ cleanup }} dashboards ({{ cleanup_pct }}%) are Terminal or Deceased and are candidates "
        "for cleanup. {{ skipped|length }} dashboards were skipped because they are provisioned or protected.",
    )
    t = doc.add_table(rows=2, cols=4)
    _borders(t)
    _widths(t, [4.25] * 4)
    for i, (label, color, key) in enumerate(
        (
            ("HEALTHY", "16A34A", "healthy"),
            ("AT RISK", "CA8A04", "at_risk"),
            ("TERMINAL", "EA580C", "terminal"),
            ("DECEASED", "DC2626", "deceased"),
        )
    ):
        top, bottom = t.rows[0].cells[i], t.rows[1].cells[i]
        _shade(top, color)
        _cell_text(top, label, bold=True, color="FFFFFF", size=9).alignment = WD_ALIGN_PARAGRAPH.CENTER
        _cell_text(bottom, "{{ counts." + key + " }}", bold=True, size=20).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    _callout(
        doc,
        "INFO",
        "This is a dry run. Gravekeeper v0.1 only reads from Grafana; nothing was changed, archived or deleted.",
        "2563EB",
    )
    _p(doc, "{%p if usage_source == 'none' %}", size=1)
    _callout(
        doc,
        "WARN",
        "View counts were not available (usage.source = none), so the Abandoned check (C01) "
        "was skipped. Ship Grafana's request logs to Loki to enable it.",
        "DC2626",
    )
    _p(doc, "{%p endif %}", size=1)

    doc.add_heading("2. Dashboards by health", 1)
    _p(
        doc,
        "Lowest score first. Score starts at 100; each cause of death subtracts its weight (see Appendix).",
        italic=True,
        color=GREY,
    )
    _loop_table(
        doc,
        ["Dashboard", "Folder", "Owner", "Views", "Score", "Status", "Causes"],
        [
            "{{ d.title }}",
            "{{ d.folder }}",
            "{{ d.owner }}",
            "{{ d.views_text }}",
            "{{ d.score }}",
            "{% cellbg d.color %}{{ d.status }}",
            "{{ d.codes }}",
        ],
        "d in dashboards",
        widths=[4.4, 2.4, 2.2, 1.4, 1.3, 1.9, 3.4],
    )

    doc.add_heading("3. Causes of death", 1)
    _loop_table(
        doc,
        ["Code", "Cause", "Dashboards", "Weight"],
        ["{{ c.code }}", "{{ c.name }}", "{{ c.count }}", "{{ c.weight }}"],
        "c in causes",
        widths=[1.6, 11.0, 2.4, 2.0],
    )

    doc.add_heading("4. Datasource health", 1)
    _loop_table(
        doc,
        ["Datasource", "Type", "UID", "Status", "Panels", "Dashboards"],
        [
            "{{ s.name }}",
            "{{ s.type }}",
            "{{ s.uid }}",
            "{% cellbg s.color %}{{ s.state }}",
            "{{ s.panels }}",
            "{{ s.dashboards }}",
        ],
        "s in datasources",
        widths=[3.6, 2.2, 2.8, 4.4, 1.6, 2.4],
    )

    doc.add_heading("5. Legacy panel migration list", 1)
    _p(doc, "{%p if legacy %}", size=1)
    _p(
        doc,
        "These Angular panels no longer render in Grafana 11/12. Opening the dashboard and saving it usually "
        "offers automatic migration; check thresholds and overrides afterwards.",
        italic=True,
        color=GREY,
    )
    _p(doc, "{%p endif %}", size=1)
    _loop_table(
        doc,
        ["Dashboard", "Panel", "Legacy type", "Replace with"],
        ["{{ l.dashboard }}", "{{ l.panel }}", "{{ l.type }}", "{{ l.replacement }}"],
        "l in legacy",
        widths=[5.0, 5.0, 3.5, 3.5],
    )
    _p(doc, "{% if not legacy %}No legacy panels found.{% endif %}")

    doc.add_heading("6. Duplicate clusters", 1)
    _loop_table(
        doc,
        ["Cluster", "Keep", "Candidates for removal"],
        ["{{ k.id }}", "{{ k.keep }}", "{{ k.others|join(', ') }}"],
        "k in duplicate_clusters",
        widths=[2.0, 5.5, 9.5],
    )
    _p(doc, "{% if not duplicate_clusters %}No duplicates found.{% endif %}")

    doc.add_heading("7. Skipped dashboards", 1)
    _loop_table(
        doc,
        ["Dashboard", "Folder", "Reason"],
        ["{{ s.title }}", "{{ s.folder }}", "{{ s.reason }}"],
        "s in skipped",
        widths=[6.0, 4.0, 7.0],
    )
    _p(doc, "{% if not skipped %}Nothing was skipped.{% endif %}")

    doc.add_heading("Appendix: scan settings", 1)
    _kv_table(
        doc,
        [
            ("Usage source", "{{ usage_source }}"),
            ("Abandoned after", "{{ config.thresholds.abandoned_days }} days without views"),
            ("Stale after", "{{ config.thresholds.stale_days }} days without edits"),
            ("Minimum age", "{{ config.thresholds.min_age_days }} days before a dashboard can be Abandoned"),
            (
                "Status bands",
                "Healthy ≥ {{ config.status_bands.healthy }}, At Risk ≥ {{ config.status_bands.at_risk }}, "
                "Terminal ≥ {{ config.status_bands.terminal }}, otherwise Deceased",
            ),
            (
                "Grace periods",
                "Terminal {{ config.grace_days.terminal }} days, Deceased {{ config.grace_days.deceased }} days",
            ),
            ("Protected tags", "{{ config.protect_tags|join(', ') }}"),
        ],
    )
    doc.save(path)


def build_death_certificate(path: Path) -> None:
    doc = _base_document()
    _banner(
        doc,
        "GRAVEKEEPER  ·  OFFICIAL NOTICE",
        "Certificate of Dashboard Death",
        "{{ d.title }}  ·  issued {{ scanned_date }}",
    )

    _p(doc, "To the owner of this dashboard ({{ d.owner }}),", bold=True)
    _p(
        doc,
        "During a routine estate scan of {{ grafana_url }}, the dashboard below was found to be "
        "{{ d.status|lower }} (health score {{ d.score }}/100). Unless you object, it will be archived on "
        "{{ burial_date }}. Archiving is reversible: the dashboard JSON is kept and can be restored at any time.",
    )

    doc.add_heading("1. Identity of the deceased", 1)
    _kv_table(
        doc,
        [
            ("Dashboard", "{{ d.title }}"),
            ("UID", "{{ d.uid }}"),
            ("Folder", "{{ d.folder }}"),
            ("Link", "{{ d.url }}"),
            ("Owner", "{{ d.owner }}"),
            ("Created", "{{ d.created_short }} by {{ d.created_by or 'unknown' }}"),
            ("Last edited", "{{ d.updated_short }} by {{ d.updated_by or 'unknown' }} ({{ d.idle_days }} days ago)"),
            ("Views", "{{ d.views_text }}"),
            ("Health score", "{{ d.score }} / 100"),
            ("Status", "{% cellbg d.color %}{{ d.status|upper }}"),
        ],
    )

    doc.add_heading("2. Cause of death", 1)
    _callout(doc, "{{ primary.code }}", "{{ primary.name }}. {{ primary.detail }}.", "DC2626")
    _p(doc, "{%p if contributing %}", size=1)
    doc.add_heading("Contributing factors", 2)
    _p(doc, "{%p endif %}", size=1)
    _loop_table(
        doc,
        ["Code", "Cause", "Detail"],
        ["{{ f.code }}", "{{ f.name }}", "{{ f.detail }}"],
        "f in contributing",
        widths=[1.6, 6.4, 9.0],
    )

    doc.add_heading("3. Evidence", 1)
    _loop_table(
        doc,
        ["Panel", "Type", "Datasource", "Finding"],
        ["{{ e.panel }}", "{{ e.type }}", "{{ e.datasource }}", "{{ e.result }}"],
        "e in d.evidence",
        widths=[4.4, 2.4, 3.2, 7.0],
    )
    _p(doc, "{% if not d.evidence %}No panel-level evidence; see the cause of death above.{% endif %}")

    doc.add_heading("4. Date of burial", 1)
    _callout(
        doc,
        "WARN",
        "Scheduled archive date: {{ burial_date }} ({{ grace_days }}-day grace period). "
        "No action is taken before this date.",
        "EA580C",
    )

    doc.add_heading("5. How to appeal", 1)
    _p(doc, "If this dashboard is still needed, do one of the following before {{ burial_date }}:")
    for step in (
        "Add the tag keep-alive to the dashboard (Dashboard settings → Tags). Protected dashboards are never archived.",
        "Add an owner tag such as owner:your-team so future notices reach the right people.",
        "Fix the causes listed above, then reply to the platform team to request a re-scan.",
    ):
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("6. Keep a copy", 1)
    _p(doc, "To export the dashboard JSON yourself right now:")
    _code(
        doc,
        'curl -s -H "Authorization: Bearer $GRAFANA_TOKEN" \\\n'
        "  {{ grafana_url }}/api/dashboards/uid/{{ d.uid }} > {{ d.uid }}.json",
    )

    doc.add_heading("Signatures", 1)
    _kv_table(
        doc,
        [
            ("Issued by", "{{ tool }} on {{ scanned_date }}"),
            ("Reviewed by", "______________________________"),
            ("Approved by", "______________________________"),
        ],
    )
    doc.save(path)


def main(out_dir: Path = OUT_DIR) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = [out_dir / "estate-report.docx", out_dir / "death-certificate.docx"]
    build_estate_report(paths[0])
    build_death_certificate(paths[1])
    return paths


if __name__ == "__main__":
    for p in main():
        print("wrote", p)
